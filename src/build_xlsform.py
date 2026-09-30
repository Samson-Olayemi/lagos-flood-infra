"""Build the digital field form (XLSForm) from docs/data_dictionary.csv.

The data dictionary is the single source of truth. The form is generated from it,
so the form and the validator in src/clean.py can never drift apart.

Input : docs/data_dictionary.csv
Output: configs/field_form_v1_0.xlsx  (XLSForm: sheets survey, choices, settings)

The XLSForm file works in KoboToolbox and ODK Collect (both free, both work offline).

Usage:
    python src/build_xlsform.py
    python src/build_xlsform.py --dictionary docs/data_dictionary.csv --out configs/field_form_v1_0.xlsx

Changing the dictionary means a new form version. Change form_version in the
dictionary (field 'form_version', column 'calculation') and FORM_VERSION below
together, then rebuild.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

FORM_VERSION = "1.0"
FORM_ID = "lagos_fi_v1_0"
FORM_TITLE = "Lagos Flood and Infrastructure Field Form v1.0 (DRAFT for pilot)"

# PRD section 19 missing-data reason codes.
MISSING_REASONS = [
    ("inaccessible", "Inaccessible"),
    ("not_applicable", "Not applicable"),
    ("device_failure", "Device failure"),
    ("obscured", "Obscured"),
    ("refused_unsafe", "Refused / unsafe"),
    ("unknown", "Unknown"),
]
MISSING_LIST = "ch_missing_reason"

SECTION_TITLES = {
    "0_site_id_and_location": "0. Site ID and location",
    "1_road": "1. Road",
    "2_drainage": "2. Drainage",
    "3_flood_evidence": "3. Flood evidence",
    "4_photographs": "4. Photographs",
    "5_notes": "5. Notes",
}


def parse_allowed(allowed: str) -> list[tuple[str, str]]:
    """'a=Label A|b=Label B' -> [('a','Label A'), ('b','Label B')]"""
    if not allowed or pd.isna(allowed):
        return []
    out = []
    for part in str(allowed).split("|"):
        key, _, label = part.partition("=")
        out.append((key.strip(), label.strip() or key.strip()))
    return out


def rule_to_xpath(rule: str) -> str:
    """Convert the dictionary's simple 'field==value' / 'field!=value' to XLSForm syntax."""
    if not rule or pd.isna(rule):
        return ""
    m = re.fullmatch(r"\s*([a-z0-9_]+)\s*(==|!=)\s*([A-Za-z0-9_]+)\s*", rule)
    if not m:
        raise ValueError(f"Unsupported relevant_if rule: {rule!r}")
    field, op, value = m.groups()
    xop = "=" if op == "==" else "!="
    # Numeric literal stays a number; words become quoted strings.
    literal = value if re.fullmatch(r"[0-9]+", value) else f"'{value}'"
    return f"${{{field}}} {xop} {literal}"


def constraint_for(row: pd.Series) -> tuple[str, str]:
    dtype = row["data_type"]
    if dtype in ("integer", "decimal"):
        parts = []
        if str(row["hard_min"]).strip() not in ("", "nan"):
            parts.append(f". >= {row['hard_min']}")
        if str(row["hard_max"]).strip() not in ("", "nan"):
            parts.append(f". <= {row['hard_max']}")
        if parts:
            unit = f" {row['unit']}" if str(row["unit"]).strip() not in ("", "nan") else ""
            msg = f"Value must be between {row['hard_min']} and {row['hard_max']}{unit}. Check the measurement."
            return " and ".join(parts), msg
    if dtype == "string" and str(row["regex"]).strip() not in ("", "nan"):
        return f"regex(., '{row['regex']}')", "Wrong format. See the hint under the question."
    return "", ""


def xlsform_type(row: pd.Series) -> tuple[str, str]:
    """Return (type, appearance) for the survey sheet."""
    src, dtype = row["source"], row["data_type"]
    if src == "auto_start":
        return "start", ""
    if src == "auto_end":
        return "end", ""
    if src == "calculated":
        return "calculate", ""
    if dtype == "geopoint":
        return "geopoint", ""
    if dtype == "photo":
        return "image", "new"
    if dtype in ("categorical", "ordinal", "binary"):
        n = len(parse_allowed(row["allowed_values"]))
        return f"select_one ch_{row['field_name']}", ("minimal" if n > 6 else "")
    if dtype == "text":
        return "text", "multiline"
    if dtype in ("string", "integer", "decimal", "date"):
        return dtype if dtype != "string" else "text", ""
    raise ValueError(f"Unsupported data_type {dtype!r} for {row['field_name']}")


def build(dictionary_path: Path, out_path: Path) -> None:
    d = pd.read_csv(dictionary_path, dtype=str, keep_default_na=False)
    wb = Workbook()

    # ---------- survey ----------
    survey = wb.active
    survey.title = "survey"
    survey_cols = ["type", "name", "label", "hint", "required", "relevant",
                   "constraint", "constraint_message", "calculation", "default", "appearance"]
    survey.append(survey_cols)

    def add_row(**kw):
        survey.append([kw.get(c, "") for c in survey_cols])

    choices_rows: list[tuple[str, str, str]] = []

    # Metadata fields go first, outside any group.
    for _, r in d[d["source"].isin(["auto_start", "auto_end"])].iterrows():
        t, _ = xlsform_type(r)
        add_row(type=t, name=r["field_name"])

    body = d[~d["source"].isin(["auto_start", "auto_end"])]
    for section, block in body.groupby("section", sort=True):
        add_row(type="begin_group", name="sec_" + re.sub(r"[^a-z0-9]+", "_", section),
                label=SECTION_TITLES.get(section, section))
        for _, r in block.iterrows():
            t, appearance = xlsform_type(r)
            relevant = rule_to_xpath(r["relevant_if"])
            constraint, cmsg = constraint_for(r)
            label = r["label"]
            if r["unit"].strip() and r["data_type"] in ("integer", "decimal"):
                label = f"{label} ({r['unit']})"
            add_row(
                type=t, name=r["field_name"], label=label, hint=r["form_hint"],
                required="yes" if (r["requirement"] == "always" and r["source"] == "form") else "no",
                relevant=relevant, constraint=constraint, constraint_message=cmsg,
                calculation=r["calculation"], default=r["default_value"], appearance=appearance,
            )
            if t.startswith("select_one"):
                for key, lab in parse_allowed(r["allowed_values"]):
                    choices_rows.append((f"ch_{r['field_name']}", key, lab))

            # Missing-reason companion: appears only when the main answer is blank.
            if r["requirement"] == "value_or_reason":
                blank = f"${{{r['field_name']}}} = ''"
                comp_rel = f"({relevant}) and {blank}" if relevant else blank
                add_row(
                    type=f"select_one {MISSING_LIST}", name=f"{r['field_name']}_missing",
                    label=f"{r['label']}: why is there no value?",
                    hint="Never leave a value blank without a reason (PRD 19).",
                    required="yes", relevant=comp_rel,
                )
        add_row(type="end_group")

    # ---------- choices ----------
    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label"])
    for lst, key, lab in choices_rows:
        choices.append([lst, key, lab])
    for key, lab in MISSING_REASONS:
        choices.append([MISSING_LIST, key, lab])

    # ---------- settings ----------
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version", "instance_name"])
    settings.append([FORM_TITLE, FORM_ID, FORM_VERSION, "${site_id}"])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Wrote {out_path}  ({survey.max_row - 1} survey rows, {len(choices_rows) + len(MISSING_REASONS)} choices)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dictionary", default="docs/data_dictionary.csv", type=Path)
    ap.add_argument("--out", default="configs/field_form_v1_0.xlsx", type=Path)
    args = ap.parse_args()
    build(args.dictionary, args.out)
